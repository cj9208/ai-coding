from __future__ import annotations

from pathlib import Path

from .base import ExtractionResult, Extractor


class DocxExtractor(Extractor):
    extensions = frozenset({"docx"})

    def extract(self, path: Path) -> ExtractionResult:
        try:
            import docx  # python-docx
        except ImportError:
            return ExtractionResult()
        document = docx.Document(str(path))
        title = document.core_properties.title or None
        text = "\n".join(p.text for p in document.paragraphs)
        return ExtractionResult(title=title, text=text or None)


class XlsxExtractor(Extractor):
    extensions = frozenset({"xlsx"})

    def extract(self, path: Path) -> ExtractionResult:
        try:
            import openpyxl  # 项目已有依赖
        except ImportError:
            return ExtractionResult()
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            parts: list[str] = []
            extra: dict = {"sheets": ",".join(wb.sheetnames)}
            for name in wb.sheetnames[:5]:  # 只扫前 5 个 sheet，控制耗时
                ws = wb[name]
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 200:  # 每个 sheet 最多 200 行
                        break
                    parts.append("\t".join("" if c is None else str(c) for c in row))
            return ExtractionResult(text="\n".join(parts) or None, extra=extra)
        finally:
            wb.close()
